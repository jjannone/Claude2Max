#!/usr/bin/env python3
"""
build_package_catalog.py — Build the package catalog spreadsheet.

Joins two files:
  - packages/package_objects.json  what each object is (ports, kind, use_when)
  - packages/package_catalog.json  judgments per object: function, category,
                                   usefulness 1-5, built-in it wraps, built-in
                                   alternative, superseded by, comments
and adds package author and version, read from each package-info.json.

Only packages installed in ~/Documents/Max 9/Packages get rows. Library
packages that are not installed are reported, not written. Compiled
externals with no Apple Silicon code are left out and reported, and `--check`
also reports names that clash with Max's own objects or another package.

Usage:
    python3 packages/build_package_catalog.py            # write packages/package_catalog.xlsx
    python3 packages/build_package_catalog.py --check    # report gaps and invalid entries only
"""

import argparse
import datetime
import json
import struct
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
LIBRARY = HERE / "package_objects.json"
CATALOG = HERE / "package_catalog.json"
OUTPUT = HERE / "package_catalog.xlsx"
USER_PACKAGES = Path.home() / "Documents" / "Max 9" / "Packages"
MAX_C74 = Path("/Applications/Max.app/Contents/Resources/C74")

CATEGORY_FILE = HERE / "package_categories.json"
# Categories and their one-line descriptions, in the order the Categories sheet lists them.
CATEGORIES = [(c["name"], c["description"]) for c in json.loads(CATEGORY_FILE.read_text())]
CATEGORY_NAMES = [name for name, _ in CATEGORIES]

USEFULNESS = {
    5: ("Everyday", "Useful in many kinds of patches, whatever the domain."),
    4: ("Broad", "A standard tool inside a big domain (audio effects, video, MIDI)."),
    3: ("Domain", "Valuable for a specific area of work (spectral, ambisonics, notation, ML)."),
    2: ("Specialized", "A narrow technique, one variant among siblings, or a part that needs its package's other objects."),
    1: ("Niche", "Very specific use, internal helper, demo, legacy, deprecated, or replaced."),
}

JUDGMENT_FIELDS = ("function", "category", "usefulness", "wraps_builtin",
                   "builtin_alternative", "superseded_by", "comments")

SOURCE_LABELS = {"refpage": "refpage", "helpfile": "help file", "abstraction": "abstraction patch"}

CPU_X86_64 = 0x01000007
CPU_ARM64 = 0x0100000C


def macho_archs(data):
    """CPU architectures in a Mach-O binary, from its header bytes."""
    if len(data) < 8:
        return set()
    magic_be = struct.unpack(">I", data[:4])[0]
    if magic_be in (0xCAFEBABE, 0xCAFEBABF):
        count = struct.unpack(">I", data[4:8])[0]
        size = 20 if magic_be == 0xCAFEBABE else 32
        archs = set()
        for i in range(count):
            off = 8 + i * size
            if len(data) < off + 4:
                break
            archs.add(struct.unpack(">I", data[off:off + 4])[0])
        return archs
    magic_le = struct.unpack("<I", data[:4])[0]
    if magic_le in (0xFEEDFACF, 0xFEEDFACE):
        return {struct.unpack("<I", data[4:8])[0]}
    return set()


def mxo_bundles(package_dir):
    """{object name: .mxo bundle path} for every compiled macOS external in a package."""
    index = {}
    for bundle in sorted(package_dir.rglob("*.mxo")):
        if bundle.is_dir():
            index.setdefault(bundle.stem, bundle)
    return index


def apple_silicon(bundle):
    """'Yes', 'No (Intel only)', 'unknown', or '' when there is no .mxo bundle."""
    if bundle is None:
        return ""
    macos = bundle / "Contents" / "MacOS"
    binaries = sorted(macos.iterdir()) if macos.is_dir() else []
    if not binaries:
        return "unknown"
    with binaries[0].open("rb") as f:
        archs = macho_archs(f.read(4096))
    if CPU_ARM64 in archs:
        return "Yes"
    if CPU_X86_64 in archs:
        return "No (Intel only)"
    return "unknown"


def max_object_names(c74=MAX_C74):
    """Every object name and alias in Max's own registries, core and bundled."""
    names = set()
    for path in [c74 / "interfaces" / "obj-qlookup.json",
                 *sorted((c74 / "packages").glob("*/interfaces/obj-qlookup.json"))]:
        try:
            data = json.loads(path.read_text())
        except (OSError, json.JSONDecodeError):
            continue
        for key, entry in data.items():
            names.add(key)
            if isinstance(entry, dict) and entry.get("alias"):
                names.add(entry["alias"])
    return names


def name_clashes(objects_by_package, builtin_names):
    """{(package, object): text} for names shared with Max or with another installed package.

    Compared without regard to case: Max on macOS finds an abstraction file
    whatever the case of the name typed in the box."""
    builtin = {}
    for name in builtin_names:
        builtin.setdefault(name.lower(), name)
    owners = {}
    for pkg, names in objects_by_package.items():
        for name in names:
            owners.setdefault(name.lower(), []).append((pkg, name))
    clashes = {}
    for key, entries in owners.items():
        for pkg, name in entries:
            parts = []
            if key in builtin:
                parts.append(f"same name as built-in {builtin[key]}")
            others = [f"{n} in {p}" for p, n in entries if p != pkg]
            if others:
                parts.append("also " + ", ".join(others))
            if parts:
                clashes[(pkg, name)] = "; ".join(parts)
    return clashes


def package_info(package_dir):
    try:
        info = json.loads((package_dir / "package-info.json").read_text(errors="replace"), strict=False)
    except (OSError, json.JSONDecodeError):
        return {}
    return info if isinstance(info, dict) else {}


def validate_entry(entry, category_names=None):
    """Problems with one catalog entry, as strings. Empty list means valid.

    `crossref` is optional: other categories the object clearly also belongs
    to, each a known category and none equal to its own."""
    names = CATEGORY_NAMES if category_names is None else category_names
    problems = []
    missing = [f for f in JUDGMENT_FIELDS if f not in entry]
    if missing:
        problems.append("missing " + ", ".join(missing))
    if entry.get("category") not in names:
        problems.append(f"unknown category {entry.get('category')!r}")
    if entry.get("usefulness") not in USEFULNESS:
        problems.append(f"usefulness {entry.get('usefulness')!r} is not 1-5")
    crossref = entry.get("crossref", [])
    if not isinstance(crossref, list):
        problems.append("crossref is not a list")
    else:
        problems += [f"unknown crossref category {c!r}" for c in crossref if c not in names]
        if entry.get("category") in crossref:
            problems.append("crossref repeats the object's own category")
    return problems


def check(library, catalog, installed, category_names=None):
    problems = []
    for pkg in sorted(installed):
        for name in sorted(library.get(pkg, {})):
            entry = catalog.get(pkg, {}).get(name)
            if entry is None:
                problems.append(f"{pkg} / {name}: no catalog entry")
            else:
                problems.extend(f"{pkg} / {name}: {p}" for p in validate_entry(entry, category_names))
    for pkg, entries in catalog.items():
        for name in entries:
            if name not in library.get(pkg, {}):
                problems.append(f"{pkg} / {name}: catalog entry for an object not in the library")
    return problems


def status_marks(record, entry):
    """The function-column suffix and the Superseded-by cell for one object.

    Deprecated comes from the library's `deprecated_by` key (the package says
    so); superseded from the catalog's `superseded_by` (a named replacement)."""
    deprecated = "deprecated_by" in record
    replacement = entry.get("superseded_by") or record.get("deprecated_by", "")
    marks = [word for word, on in (("DEPRECATED", deprecated), ("SUPERSEDED", bool(entry.get("superseded_by")))) if on]
    return " ".join(marks), replacement or ("DEPRECATED" if deprecated else "")


def build_rows(library, catalog, installed):
    """Rows for the Objects sheet, and the (package, object) pairs left out as Intel-only."""
    rows, intel_only = [], []
    for pkg in sorted(installed, key=str.lower):
        bundles = mxo_bundles(USER_PACKAGES / pkg)
        for name in sorted(library[pkg], key=str.lower):
            if apple_silicon(bundles.get(name)) == "No (Intel only)":
                intel_only.append((pkg, name))
                continue
            rec = library[pkg][name]
            j = catalog.get(pkg, {}).get(name, {})
            level = USEFULNESS.get(j.get("usefulness"), ("", ""))[0]
            unknown_ports = not rec["numinlets"] and not rec["numoutlets"]
            marks, superseded = status_marks(rec, j)
            function = j.get("function", "")
            rows.append({
                "Object": name, "Package": pkg,
                "General function": f"{function} {marks}".strip(), "Category": j.get("category", ""),
                "Also in": "; ".join(j.get("crossref", [])),
                "Usefulness (1-5)": j.get("usefulness"), "Usefulness level": level,
                "Kind": rec["kind"],
                "Inlets": None if unknown_ports else rec["numinlets"],
                "Outlets": None if unknown_ports else rec["numoutlets"],
                "Wraps built-in": j.get("wraps_builtin", ""),
                "Built-in alternative": j.get("builtin_alternative", ""),
                "Superseded by": superseded,
                "use_when": rec["use_when"],
                "Doc source": SOURCE_LABELS.get(rec["source"], rec["source"]),
                "Comments": j.get("comments", ""),
            })
    # Row order as John sorted it in Excel (2026-09-16): category A-Z, most useful first.
    rows.sort(key=lambda r: (r["Category"].lower(), -(r["Usefulness (1-5)"] or 0),
                             r["Package"].lower(), r["Object"].lower()))
    return rows, intel_only


# Order and widths as John arranged them in Excel (2026-09-16).
COLUMNS = [
    ("Object", 26, "The name you type in an object box."),
    ("Package", 9, "The package folder in ~/Documents/Max 9/Packages."),
    ("General function", 44, "What it does, in one plain sentence."),
    ("Category", 26, "One primary category; see the Categories sheet."),
    ("Also in", 26, "Other categories this object clearly belongs to as well, for instance a random-walk melody generator filed under MIDI that is also a random generator."),
    ("Usefulness (1-5)", 11, "5 = everyday patching, 1 = highly specific use. See the Usefulness sheet."),
    ("Usefulness level", 15.66, "The name of the usefulness rating."),
    ("Kind", 15.66, "external (compiled), abstraction (a Max patch), or javascript."),
    ("Built-in alternative", 34, "A built-in Max object that does the same or a similar job, and how it differs."),
    ("Inlets", 7, "Inlet count. Blank when unknown (for instance ports built by a script)."),
    ("Outlets", 7, "Outlet count. Blank when unknown."),
    ("Comments", 44, "Gotchas, required companions, platform limits, anything else that changes a decision."),
    ("Wraps built-in", 26.33, "The built-in Max object(s) doing the real work, when this is essentially a wrapper or front panel for them. Built-in includes Max's bundled packages."),
    ("Superseded by", 20, "A newer object that replaces this one, or DEPRECATED when the package deprecates it without naming a replacement. General function ends in SUPERSEDED and/or DEPRECATED to match."),
    ("use_when", 60, "The curated guidance from packages/package_objects.json: when to reach for it and how to drive it."),
    ("Doc source", 12, "Where the object's documentation came from: refpage, help file, or the abstraction patch itself."),
]
COLUMN_NAMES = [name for name, _, _ in COLUMNS]

LEVEL_FILLS = {5: "C6EFCE", 4: "DDEBF7", 3: "FFF2CC", 2: "F2F2F2", 1: "E7E6E6"}


def write_workbook(rows, library, catalog, installed, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    font = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    wrap = Alignment(vertical="top", wrap_text=True)

    def header(ws, names_widths):
        for col, (name, width) in enumerate(names_widths, 1):
            cell = ws.cell(row=1, column=col, value=name)
            cell.font, cell.fill, cell.alignment = bold, header_fill, wrap
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = "B2"

    # Formatting as John set it in Excel (2026-09-17): 16 pt object names,
    # 35 pt rows, numeric columns centred.
    big, big_bold = Font(name="Arial", size=16), Font(name="Arial", size=16, bold=True)
    centred = {"Usefulness (1-5)", "Inlets", "Outlets"}
    wrapped = {"General function", "Also in", "Wraps built-in", "Built-in alternative", "use_when", "Comments"}

    def align(name, is_header=False):
        return Alignment(horizontal="center" if name in centred else None, vertical="top",
                         wrap_text=True if is_header or name in wrapped else None)

    wb = Workbook()
    ws = wb.active
    ws.title = "Objects"
    header(ws, [(n, w) for n, w, _ in COLUMNS])
    for c, name in enumerate(COLUMN_NAMES, 1):
        cell = ws.cell(row=1, column=c)
        cell.alignment = align(name, is_header=True)
        if name == "Object":
            cell.font = big_bold
    ws.sheet_format.defaultRowHeight = 35
    ws.sheet_format.customHeight = True
    ws.row_dimensions[1].height = 35
    level_cols = [COLUMN_NAMES.index("Usefulness (1-5)") + 1, COLUMN_NAMES.index("Usefulness level") + 1]
    for r, row in enumerate(rows, 2):
        ws.row_dimensions[r].height = 35
        for c, name in enumerate(COLUMN_NAMES, 1):
            value = row[name]
            cell = ws.cell(row=r, column=c, value=value if value != "" else None)
            cell.font = big if name == "Object" else font
            cell.alignment = align(name)
        if row["Usefulness (1-5)"] in LEVEL_FILLS:
            fill = PatternFill("solid", fgColor=LEVEL_FILLS[row["Usefulness (1-5)"]])
            for c in level_cols:
                ws.cell(row=r, column=c).fill = fill
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"

    counted = f"Counted by packages/build_package_catalog.py on {datetime.date.today().isoformat()}."

    wp = wb.create_sheet("Packages")
    header(wp, [("Package", 30), ("Author", 28), ("Version", 10), ("Objects", 9),
                ("Main category", 28), ("Everyday or broad (4-5)", 12), ("Description", 80)])
    for r, pkg in enumerate(sorted(installed, key=str.lower), 2):
        info = package_info(USER_PACKAGES / pkg)
        pkg_rows = [row for row in rows if row["Package"] == pkg]
        cats = [row["Category"] for row in pkg_rows if row["Category"]]
        main = max(sorted(set(cats), key=CATEGORY_NAMES.index), key=cats.count) if cats else ""
        text = lambda v: ", ".join(map(str, v)) if isinstance(v, list) else str(v or "")
        values = [pkg, text(info.get("author")), text(info.get("version")), len(pkg_rows), main,
                  sum(1 for row in pkg_rows if (row["Usefulness (1-5)"] or 0) >= 4), text(info.get("description"))]
        for c, value in enumerate(values, 1):
            cell = wp.cell(row=r, column=c, value=value if value != "" else None)
            cell.font, cell.alignment = font, wrap
    wp.cell(row=len(installed) + 3, column=1, value=counted).font = font

    wc = wb.create_sheet("Categories")
    header(wc, [("Category", 34), ("What belongs here", 70), ("Objects", 9)])
    for r, (name, desc) in enumerate(CATEGORIES, 2):
        for c, value in enumerate([name, desc, sum(1 for row in rows if row["Category"] == name)], 1):
            cell = wc.cell(row=r, column=c, value=value)
            cell.font, cell.alignment = font, wrap
    wc.cell(row=len(CATEGORIES) + 3, column=1, value=counted).font = font

    wu = wb.create_sheet("Usefulness")
    header(wu, [("Rating", 8), ("Level", 14), ("Meaning", 80), ("Objects", 9)])
    for r, rating in enumerate(sorted(USEFULNESS, reverse=True), 2):
        level, meaning = USEFULNESS[rating]
        for c, value in enumerate([rating, level, meaning, sum(1 for row in rows if row["Usefulness (1-5)"] == rating)], 1):
            cell = wu.cell(row=r, column=c, value=value)
            cell.font, cell.alignment = font, wrap
        wu.cell(row=r, column=1).fill = PatternFill("solid", fgColor=LEVEL_FILLS[rating])
    wu.cell(row=8, column=1, value="Ratings compare each object with what Max already does: "
            "an object a built-in does just as easily is rated lower.").font = font
    wu.cell(row=9, column=1, value=counted).font = font

    wd = wb.create_sheet("Columns")
    header(wd, [("Column", 24), ("Meaning", 100)])
    for r, (name, _, meaning) in enumerate(COLUMNS, 2):
        for c, value in enumerate([name, meaning], 1):
            cell = wd.cell(row=r, column=c, value=value)
            cell.font, cell.alignment = font, wrap
    wd.cell(row=len(COLUMNS) + 3, column=1,
            value="Sources: packages/package_objects.json and packages/package_catalog.json in the "
                  "Claude2Max repo. Rebuild with python3 packages/build_package_catalog.py.").font = font

    wb.save(path)


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--check", action="store_true", help="report problems without writing")
    p.add_argument("-o", "--output", default=str(OUTPUT))
    args = p.parse_args()

    library = json.loads(LIBRARY.read_text())
    catalog = json.loads(CATALOG.read_text()) if CATALOG.exists() else {}
    if not USER_PACKAGES.is_dir():
        sys.exit(f"Packages folder not found: {USER_PACKAGES}")
    present = {d.name for d in USER_PACKAGES.iterdir() if d.is_dir()}
    installed = sorted(p for p in library if p in present)
    absent = sorted(p for p in library if p not in present)
    if absent:
        print(f"In the library but not installed, left out: {', '.join(absent)}", file=sys.stderr)

    problems = check(library, catalog, installed)
    for line in problems:
        print(line, file=sys.stderr)
    print(f"{len(problems)} problem(s) across {len(installed)} installed packages", file=sys.stderr)
    if args.check:
        clashes = name_clashes({p: list(library[p]) for p in installed}, max_object_names())
        for (pkg, name), text in sorted(clashes.items()):
            print(f"name clash: {pkg} / {name}: {text}", file=sys.stderr)
        sys.exit(1 if problems else 0)

    rows, intel_only = build_rows(library, catalog, installed)
    if intel_only:
        print("Intel-only externals, left out: " + ", ".join(f"{p} / {n}" for p, n in intel_only), file=sys.stderr)
    write_workbook(rows, library, catalog, installed, args.output)
    print(f"Wrote {len(rows)} objects to {args.output}", file=sys.stderr)


if __name__ == "__main__":
    main()
